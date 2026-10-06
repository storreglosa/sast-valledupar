"""Informe ejecutivo (gerencia) y Excel para copiar a la plataforma ANSV.

Lenguaje sin tecnicismos de bases de datos. Cifra oficial: zona de influencia + 15 m
(decisión 25). Todas las cifras llegan calculadas desde `scripts/08_reporte_ejecutivo.py`.
"""

from __future__ import annotations

import html
import json

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# Descripción abreviada de cada código (Código Nacional de Tránsito, art. 131, Ley 769/2002
# modificada por la Ley 1383/2010). Solo para orientar al lector; no reemplaza el texto legal.
INFRACCIONES = {
    "C02": "Estacionar en sitio prohibido",
    "C03": "Bloquear calzada o intersección",
    "C24": "Conducir motocicleta sin cumplir las normas",
    "C29": "Exceder la velocidad máxima permitida",
    "C32": "No respetar el paso de peatones",
    "C35": "No realizar la revisión técnico-mecánica",
    "D02": "Conducir sin el seguro obligatorio (SOAT)",
    "D03": "Transitar en contravía",
    "D04": "No detenerse ante semáforo en rojo o señal de PARE",
    "D05": "Conducir sobre andenes, separadores o zonas verdes",
}
MESES_ES = ["Enero", "Febrero", "Marzo", "Abril", "Mayo", "Junio", "Julio", "Agosto",
            "Septiembre", "Octubre", "Noviembre", "Diciembre"]
MESES_CORTOS = ["ene", "feb", "mar", "abr", "may", "jun", "jul", "ago", "sep", "oct", "nov", "dic"]


def esc(t) -> str:
    return html.escape(str(t))


def fmt(n) -> str:
    return f"{int(round(n)):,}".replace(",", ".")


def mes_largo(p: pd.Period) -> str:
    return f"{MESES_ES[p.month - 1]} {p.year}"


def mes_corto(p: pd.Period) -> str:
    return f"{MESES_CORTOS[p.month - 1]} {p.year}"


def rango(v: pd.PeriodIndex) -> str:
    return f"{mes_corto(v[0])} – {mes_corto(v[-1])}"


# --------------------------------------------------------------------------------------------
# Excel para la plataforma
# --------------------------------------------------------------------------------------------
NARANJA = PatternFill("solid", fgColor="E07B4C")
GRIS = PatternFill("solid", fgColor="EEF0F2")
CLARO = PatternFill("solid", fgColor="FBF3EE")
FINO = Side(style="thin", color="C9CDD1")
BORDE = Border(left=FINO, right=FINO, top=FINO, bottom=FINO)
CENTRO = Alignment(horizontal="center", vertical="center", wrap_text=True)
OBS = {"Fallecidos": "Personas fallecidas en siniestros ocurridos en la zona de influencia del equipo.",
       "Lesionados": "Personas lesionadas en siniestros ocurridos en la zona de influencia del equipo."}
OBS_COMP = "Comparendos impuestos en la zona de influencia del equipo (agentes y fotodetección previa)."


def excel_ansv(equipos: list[dict], resumen: pd.DataFrame, instrucciones: list[str]) -> Workbook:
    """equipos: [{info..., 'valores': DataFrame indicador × Period (36 meses)}]."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Instrucciones"
    ws["A1"] = "Línea base SAST Valledupar — cifras para la plataforma ANSV"
    ws["A1"].font = Font(bold=True, size=14)
    for i, t in enumerate(instrucciones, 3):
        ws.cell(i, 1, t).alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 120

    rs = wb.create_sheet("Resumen")
    cols = list(resumen.columns)
    for j, c in enumerate(cols, 1):
        cel = rs.cell(1, j, c)
        cel.font, cel.fill, cel.border, cel.alignment = Font(bold=True, color="FFFFFF"), NARANJA, BORDE, CENTRO
    for i, fila in enumerate(resumen.itertuples(index=False), 2):
        for j, v in enumerate(fila, 1):
            c = rs.cell(i, j, None if (isinstance(v, float) and pd.isna(v)) else v)
            c.border, c.alignment = BORDE, CENTRO
    for j, c in enumerate(cols, 1):
        rs.column_dimensions[get_column_letter(j)].width = max(10, min(46, len(str(c)) + 4))
    rs.column_dimensions["E"].width = 46
    rs.freeze_panes = "B2"

    for e in equipos:
        ws = wb.create_sheet(e["equipo"])
        ws["A1"] = f"{e['equipo']} · {e['punto']}"
        ws["A1"].font = Font(bold=True, size=14)
        datos = [("Código solicitud", e["solicitud_ansv"]), ("Código único", e["codigo_unico"]),
                 ("Dirección", e["direccion_ansv"]), ("Fecha inicio de operación", e["fecha_inicio"]),
                 ("Línea base", f"{mes_largo(e['ventana'][0])} a {mes_largo(e['ventana'][-1])} (36 meses)")]
        for i, (k, v) in enumerate(datos, 2):
            ws.cell(i, 1, k).font = Font(bold=True)
            ws.cell(i, 2, v)
        fila = 8
        val = e["valores"]
        for a in range(3):
            meses = list(e["ventana"][a * 12:(a + 1) * 12])
            ws.cell(fila, 1, f"Año {a + 1}").font = Font(bold=True, size=12, color="FFFFFF")
            ws.cell(fila, 1).fill = NARANJA
            ws.cell(fila, 2, f"{mes_largo(meses[0])} a {mes_largo(meses[-1])}").font = Font(italic=True)
            fila += 1
            cab = ["Indicador"] + [mes_largo(m) for m in meses] + ["Total año", "Observaciones"]
            for j, c in enumerate(cab, 1):
                cel = ws.cell(fila, j, c)
                cel.font, cel.fill, cel.border, cel.alignment = Font(bold=True), GRIS, BORDE, CENTRO
            ws.row_dimensions[fila].height = 30
            fila += 1
            for ind in val.index:
                ws.cell(fila, 1, ind).font = Font(bold=True)
                ws.cell(fila, 1).border = BORDE
                for j, m in enumerate(meses, 2):
                    c = ws.cell(fila, j, int(val.at[ind, m]))
                    c.border, c.alignment = BORDE, CENTRO
                c = ws.cell(fila, 14, int(sum(val.at[ind, m] for m in meses)))
                c.border, c.alignment, c.font, c.fill = BORDE, CENTRO, Font(bold=True), CLARO
                c = ws.cell(fila, 15, OBS.get(ind, OBS_COMP))
                c.border, c.alignment = BORDE, Alignment(wrap_text=True, vertical="center")
                fila += 1
            fila += 1
        ws.column_dimensions["A"].width = 13
        ws.column_dimensions["B"].width = 16
        for j in range(3, 14):
            ws.column_dimensions[get_column_letter(j)].width = 12
        ws.column_dimensions["N"].width = 10
        ws.column_dimensions["O"].width = 60
    return wb


# --------------------------------------------------------------------------------------------
# Informe ejecutivo HTML
# --------------------------------------------------------------------------------------------
CSS = """
/* Layout: una columna editorial con bandas a todo el ancho (cifras, mapa) y una retícula de
   fichas por punto. Acento único: el naranja de la plataforma ANSV; la línea amarilla
   discontinua de la demarcación vial es la única ornamentación. */
:root{
  --paper:#f4f5f2; --surface:#ffffff; --ink:#1c2227; --ink-2:#48525a; --ink-3:#78828a;
  --line:#dde1dc; --accent:#c95a24; --accent-soft:#f6e7de; --signal:#e5b100;
  --dead:#b8332a; --hurt:#d08a16; --comp:#2f6fb2; --s-ag:#2a78d6; --s-fo:#eb6834; --base:#f3e3d6;
  --display:"Barlow Condensed","Arial Narrow",Arial,sans-serif;
  --body:"Public Sans","Segoe UI",system-ui,sans-serif;
  color-scheme:light;
}
@media (prefers-color-scheme:dark){:root:not([data-theme="light"]){
  --paper:#13171a; --surface:#1b2024; --ink:#eef1ee; --ink-2:#b9c1c6; --ink-3:#8b959c;
  --line:#2c343a; --accent:#e5824f; --accent-soft:#3a2a22; --signal:#d9a900;
  --dead:#e0635a; --hurt:#e2a43a; --comp:#5d9be0; --s-ag:#3987e5; --s-fo:#d95926; --base:#33261e; color-scheme:dark}}
:root[data-theme="dark"]{
  --paper:#13171a; --surface:#1b2024; --ink:#eef1ee; --ink-2:#b9c1c6; --ink-3:#8b959c;
  --line:#2c343a; --accent:#e5824f; --accent-soft:#3a2a22; --signal:#d9a900;
  --dead:#e0635a; --hurt:#e2a43a; --comp:#5d9be0; --s-ag:#3987e5; --s-fo:#d95926; --base:#33261e; color-scheme:dark}
*{box-sizing:border-box}
html,body{margin:0}
body{background:var(--paper);color:var(--ink);font:16px/1.6 var(--body);-webkit-font-smoothing:antialiased}
.envoltura{max-width:1120px;margin:0 auto;padding-inline:20px;padding-block:0 72px}
h1,h2,h3{font-family:var(--display);text-wrap:balance;line-height:1.05;margin:0}
h1{font-size:clamp(2.4rem,5.5vw,4rem);font-weight:700;letter-spacing:-.01em}
h2{font-size:clamp(1.7rem,3vw,2.2rem);font-weight:600}
h3{font-size:1.45rem;font-weight:600}
p{margin:0}
.eyebrow{font:600 .78rem/1.3 var(--body);letter-spacing:.12em;text-transform:uppercase;color:var(--accent)}
.lead{font-size:1.12rem;color:var(--ink-2);max-width:62ch}
.nota{font-size:.86rem;color:var(--ink-3);max-width:75ch}
section{display:grid;gap:18px;padding-block:44px 0}
.cabecera{display:grid;gap:18px;padding-block:52px 28px}
.demarcacion{height:6px;background:repeating-linear-gradient(90deg,var(--signal) 0 46px,transparent 46px 74px);border-radius:2px;opacity:.9}
.meta{display:flex;flex-wrap:wrap;gap:10px 28px;font-size:.9rem;color:var(--ink-2)}
.meta b{color:var(--ink);font-weight:600}
.estado{display:inline-flex;align-items:center;gap:8px;background:var(--accent-soft);color:var(--accent);
  font-weight:600;font-size:.82rem;padding:4px 12px;border-radius:999px}
.estado::before{content:"";width:8px;height:8px;border-radius:50%;background:var(--accent)}
.cifras{display:grid;grid-template-columns:repeat(4,minmax(0,1fr));border-block:1px solid var(--line)}
.cifra{padding:22px 18px;display:grid;gap:4px;align-content:start}
.cifra+.cifra{border-left:1px solid var(--line)}
.cifra .n{font:700 clamp(2.4rem,5vw,3.6rem)/1 var(--display);font-variant-numeric:tabular-nums}
.cifra .r{font-weight:600}
.cifra .d{font-size:.84rem;color:var(--ink-3)}
.cifra.muertos .n{color:var(--dead)} .cifra.heridos .n{color:var(--hurt)}
@media (max-width:760px){.cifras{grid-template-columns:repeat(2,minmax(0,1fr))}
  .cifra:nth-child(3){border-left:0}.cifra:nth-child(n+3){border-top:1px solid var(--line)}}
.fila-botones{display:flex;flex-wrap:wrap;gap:8px}
.fila-botones.sub{padding-left:14px;border-left:3px solid var(--accent)}
.fila-botones.sub button{font-size:.84rem;padding:5px 12px}
.fila-botones button{font:600 .9rem var(--body);padding:7px 14px;border:1px solid var(--line);border-radius:6px;
  background:var(--surface);color:var(--ink);cursor:pointer}
.fila-botones button:hover,.fila-botones button:focus-visible{border-color:var(--accent);color:var(--accent);outline:none}
.fila-botones button[aria-pressed="true"]{background:var(--ink);color:var(--paper);border-color:var(--ink)}
#mapa{height:min(68vh,600px);min-height:380px;border:1px solid var(--line);border-radius:10px;overflow:hidden}
.leyenda{display:flex;flex-wrap:wrap;gap:8px 22px;font-size:.88rem;color:var(--ink-2)}
.leyenda span{display:inline-flex;align-items:center;gap:8px}
.sw{width:12px;height:12px;border-radius:50%;display:inline-block}
.sw.eq{border-radius:3px;background:var(--ink)} .sw.zona{border-radius:3px;background:var(--accent-soft);border:2px solid var(--accent)}
.sw.comp{background:var(--comp)} .sw.ag{background:var(--s-ag);border-radius:3px} .sw.fo{background:var(--s-fo);border-radius:3px}
.sw.mu{background:var(--dead);border-radius:3px} .sw.he{background:var(--hurt);border-radius:3px} .sw.muerto{background:var(--dead)} .sw.herido{background:var(--hurt)}
.placa-mapa span{display:block;background:#1c2227;color:#fff;font:700 12px/20px "Barlow Condensed",Arial,sans-serif;
  letter-spacing:.04em;text-align:center;border-radius:4px;border:2px solid #fff;box-shadow:0 1px 4px rgba(0,0,0,.45)}
.estado-mapa{font-size:.86rem;color:var(--ink-3)}
.tablero{background:var(--surface);border:1px solid var(--line);border-radius:14px;padding:clamp(16px,3vw,28px);display:grid;gap:22px;min-width:0}
.filtros{display:grid;gap:10px}
.t-cab{display:flex;flex-wrap:wrap;justify-content:space-between;gap:18px;align-items:end;border-bottom:1px solid var(--line);padding-bottom:18px}
.t-cab>div{min-width:0}
.t-sol{font-size:.8rem;color:var(--ink-3);letter-spacing:.06em;text-transform:uppercase}
.t-cab h3{font-size:clamp(1.6rem,3vw,2.1rem);margin-block:2px 6px}
.t-meta{font-size:.9rem;color:var(--ink-2)}
.t-kpis{display:flex;gap:28px}
.t-kpis div{display:grid}
.t-kpis .n{font:700 2.4rem/1 var(--display);font-variant-numeric:tabular-nums}
.t-kpis .r{font-size:.8rem;color:var(--ink-3)}
.t-kpis .m .n{color:var(--dead)} .t-kpis .h .n{color:var(--hurt)}
.t-bloque{display:grid;gap:8px;min-width:0}
.t-bloque h4{margin:0;font:600 .8rem var(--body);letter-spacing:.08em;text-transform:uppercase;color:var(--ink-3)}
.graf-env{position:relative;min-width:0}
.graf-env svg{width:100%;height:auto;display:block;overflow:visible}
.graf-env .grid{stroke:var(--line)} .graf-env .eje{fill:var(--ink-3);font:12px var(--body)}
.graf-env .base{stroke:var(--ink-3)} .graf-env .ventana{fill:var(--base)}
.graf-env .anio{fill:var(--accent);font:600 12px var(--body)} .graf-env .sepa{stroke:var(--accent);stroke-dasharray:4 4;opacity:.6}
.graf-env .inicio{stroke:var(--ink);stroke-width:1.5;stroke-dasharray:5 4} .graf-env .inicio-t{fill:var(--ink);font:600 12px var(--body)}
.graf-env .hit{fill:transparent;cursor:crosshair} .graf-env .col:hover .hit{fill:var(--line);opacity:.6}
.tip{position:absolute;pointer-events:none;background:var(--ink);color:var(--paper);font-size:.82rem;line-height:1.45;
  padding:8px 11px;border-radius:7px;white-space:nowrap;box-shadow:0 4px 14px rgba(0,0,0,.25);z-index:5}
.tip b{font-weight:600} .tip .sw{width:9px;height:9px;margin-right:6px}
.cods{display:grid;gap:4px}
.cod{display:grid;grid-template-columns:3.4em minmax(0,15em) minmax(60px,1fr) 4.5em;gap:12px;align-items:center;padding:7px 0;border-bottom:1px solid var(--line);font-size:.9rem}
.cod:last-child{border-bottom:0}
.cod code{font:700 1.05rem var(--display);color:var(--accent)}
.cod .desc{color:var(--ink-2);line-height:1.3}
.cod .bar{display:flex;height:12px;border-radius:3px;overflow:hidden;background:var(--paper)}
.cod .bar i{display:block;height:100%} .cod .bar .ag{background:var(--s-ag)} .cod .bar .fo{background:var(--s-fo)}
.cod .tot{text-align:right;font:700 1.15rem var(--display);font-variant-numeric:tabular-nums}
.cod .det{grid-column:2/-1;font-size:.78rem;color:var(--ink-3);margin-top:-4px}
.cod.cero{opacity:.55}
@media (max-width:640px){.cod{grid-template-columns:3em minmax(0,1fr) 4em}.cod .bar{grid-column:2/-1;grid-row:2}.cod .det{grid-row:3}}
.tabla{overflow-x:auto;border:1px solid var(--line);border-radius:10px;background:var(--surface)}
table{border-collapse:collapse;width:100%;font-variant-numeric:tabular-nums;font-size:.92rem}
th,td{padding:10px 12px;text-align:right;border-bottom:1px solid var(--line);white-space:nowrap}
th{font:600 .78rem var(--body);letter-spacing:.06em;text-transform:uppercase;color:var(--ink-3);background:var(--paper)}
th:first-child,td:first-child{text-align:left;position:sticky;left:0;background:inherit}
tbody tr{background:var(--surface)} tbody tr:last-child td{border-bottom:0}
td.eqc b{font:700 1.05rem var(--display)} td.eqc small{display:block;color:var(--ink-3);font-size:.75rem}
td.na{color:var(--ink-3)} td.vm{color:var(--dead);font-weight:600} td.vh{color:var(--hurt);font-weight:600}
.pasos{display:grid;grid-template-columns:repeat(auto-fit,minmax(min(100%,230px),1fr));gap:18px;counter-reset:paso;margin:0;padding:0;list-style:none}
.pasos li{display:grid;gap:6px;align-content:start;padding-top:14px;border-top:3px solid var(--ink);counter-increment:paso}
.pasos li::before{content:counter(paso);font:700 1.6rem/1 var(--display);color:var(--accent)}
.pasos b{font-weight:600}
.pasos p{font-size:.92rem;color:var(--ink-2)}
.consideraciones{display:grid;gap:12px;margin:0;padding:0;list-style:none;max-width:80ch}
.consideraciones li{padding-left:18px;border-left:3px solid var(--signal);color:var(--ink-2)}
.consideraciones b{color:var(--ink)}
.traza{display:grid;grid-template-columns:repeat(auto-fit,minmax(min(100%,300px),1fr));gap:18px}
.traza dl{margin:0;display:grid;gap:10px}
.traza dt{font-size:.75rem;letter-spacing:.08em;text-transform:uppercase;color:var(--ink-3)}
.traza dd{margin:0;font-size:.95rem}
.traza code{font-size:.85rem;background:var(--paper);padding:1px 6px;border-radius:4px}
.pie{margin-top:56px;padding-top:18px;border-top:1px solid var(--line);font-size:.82rem;color:var(--ink-3);display:flex;flex-wrap:wrap;gap:8px 24px}
.leaflet-tooltip{font:13px/1.4 var(--body)}
"""


def pagina(c: dict, leaflet_css: str) -> str:
    """c: diccionario con todas las cifras y textos ya calculados."""
    datos = json.dumps(c["mapa"], ensure_ascii=False, separators=(",", ":"))
    tablero = json.dumps(c["tablero"], ensure_ascii=False, separators=(",", ":"))
    cods = c["codigos_tabla"]
    filas = []
    for e in c["tabla"]:
        celdas = [f'<td class="vm">{fmt(e["Fallecidos"])}</td>', f'<td class="vh">{fmt(e["Lesionados"])}</td>']
        celdas += [f"<td>{fmt(e[k])}</td>" if k in e else '<td class="na">—</td>' for k in cods]
        filas.append(f'<tr><td class="eqc"><b>{esc(e["equipo"])}</b><small>{esc(e["punto"])} · {esc(e["ventana"])}</small></td>'
                     + "".join(celdas) + "</tr>")
    cab = "".join(f'<th title="{esc(INFRACCIONES.get(k, ""))}">{esc(k)}</th>' for k in cods)
    leyenda_cod = " · ".join(f"<b>{esc(k)}</b> {esc(INFRACCIONES[k].lower())}" for k in cods if k in INFRACCIONES)
    fuentes = "".join(f"<dt>{esc(a)}</dt><dd>{b}</dd>" for a, b in c["fuentes"])
    entregables = "".join(f"<dt>{esc(a)}</dt><dd>{b}</dd>" for a, b in c["entregables"])
    return f"""<!doctype html>
<html lang="es"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1,viewport-fit=cover">
<title>Línea base SAST Valledupar</title>
<link rel="preconnect" href="https://fonts.googleapis.com"><link rel="preconnect" href="https://fonts.gstatic.com" crossorigin>
<link rel="stylesheet" href="https://fonts.googleapis.com/css2?family=Barlow+Condensed:wght@500;600;700&family=Public+Sans:wght@400;600&display=swap">
<style>{leaflet_css}</style>
<style>{CSS}</style>
</head><body><div class="envoltura">

<header class="cabecera">
  <span class="eyebrow">Secretaría de Tránsito y Transporte de Valledupar · Reporte a la Agencia Nacional de Seguridad Vial</span>
  <h1>Línea base de seguridad vial de los equipos de fotodetección</h1>
  <p class="lead">Estado de la seguridad vial en la zona de influencia de cada equipo de detección electrónica (SAST) durante los
    36 meses previos a su entrada en operación: personas fallecidas y lesionadas en siniestros, y comparendos impuestos por
    cada infracción que el equipo controla. Es el punto de partida con el que se medirá su efecto.</p>
  <div class="meta"><span class="estado">Lista para reporte a la ANSV</span>
    <span>Comparendos con corte al <b>{esc(c["corte_comparendos"])}</b></span>
    <span>Siniestros con corte al <b>{esc(c["corte_siniestros"])}</b></span>
    <span>Validado por <b>{esc(c["validado"])}</b></span></div>
  <div class="demarcacion" aria-hidden="true"></div>
</header>

<div class="cifras">
  <div class="cifra"><span class="n">{c["n_equipos"]}</span><span class="r">equipos en operación</span><span class="d">en {c["n_puntos"]} puntos de la ciudad</span></div>
  <div class="cifra"><span class="n">{fmt(c["tot_comparendos"])}</span><span class="r">comparendos</span><span class="d">por las infracciones que controlan los equipos</span></div>
  <div class="cifra muertos"><span class="n">{fmt(c["tot_fallecidos"])}</span><span class="r">personas fallecidas</span><span class="d">en siniestros dentro de las zonas</span></div>
  <div class="cifra heridos"><span class="n">{fmt(c["tot_lesionados"])}</span><span class="r">personas lesionadas</span><span class="d">en siniestros dentro de las zonas</span></div>
</div>
<p class="nota" style="margin-top:10px">Totales de los {c["n_puntos"]} puntos en sus 36 meses de línea base. Donde las zonas de dos equipos de un mismo punto se
  superponen, cada hecho se cuenta una sola vez.</p>

<section id="mapa-seccion">
  <div><span class="eyebrow">Mapa interactivo</span><h2>Dónde están los equipos y qué pasó a su alrededor</h2></div>
  <p class="lead">Cada placa negra es un equipo. El área naranja es su zona de influencia. Los puntos muestran los comparendos y los
    siniestros con víctimas de la línea base. Elija un punto y luego uno de sus equipos para ver solo su zona; acerque o aleje con la
    rueda del mouse y pase el cursor sobre los elementos para ver el detalle.</p>
  <div class="filtros"><div id="botones" class="fila-botones" role="group" aria-label="Filtrar el mapa por punto"></div>
    <div id="botones-eq" class="fila-botones sub" role="group" aria-label="Filtrar el mapa por equipo" hidden></div></div>
  <div id="mapa" aria-label="Mapa de equipos SAST, zonas de influencia y hechos de la línea base"></div>
  <div class="leyenda"><span><i class="sw eq"></i>Equipo SAST</span><span><i class="sw zona"></i>Zona de influencia</span>
    <span><i class="sw comp"></i>Comparendo</span><span><i class="sw muerto"></i>Siniestro con fallecidos</span>
    <span><i class="sw herido"></i>Siniestro con lesionados</span></div>
  <p class="estado-mapa" id="estado-mapa"></p>
</section>

<section id="tablero-sec">
  <div><span class="eyebrow">Tablero por punto y equipo</span><h2>Comportamiento mes a mes</h2></div>
  <p class="lead">Elija un punto y, si quiere, uno de sus equipos. La franja sombreada son los 36 meses de la línea base, dividida en
    sus tres años; la línea punteada marca el inicio de operación. Pase el cursor sobre cada mes para ver el detalle.</p>
  <div class="tablero">
    <div class="filtros"><div id="t-puntos" class="fila-botones" role="group" aria-label="Elegir punto"></div>
      <div id="t-equipos" class="fila-botones sub" role="group" aria-label="Elegir equipo"></div></div>
    <div class="t-cab"><div><span class="t-sol" id="t-sol"></span><h3 id="t-nombre"></h3><p class="t-meta" id="t-meta"></p></div>
      <div class="t-kpis"><div><span class="n" id="k-comp"></span><span class="r">comparendos en la línea base</span></div>
        <div class="m"><span class="n" id="k-fall"></span><span class="r">fallecidos</span></div>
        <div class="h"><span class="n" id="k-les"></span><span class="r">lesionados</span></div></div></div>
    <div class="t-bloque"><h4>Comparendos por mes</h4>
      <div class="leyenda"><span><i class="sw ag"></i>Agentes de tránsito</span><span><i class="sw fo"></i>Fotodetección anterior a los equipos SAST</span></div>
      <div class="graf-env" id="g-comp"></div></div>
    <div class="t-bloque"><h4>Personas fallecidas y lesionadas por mes</h4>
      <div class="leyenda"><span><i class="sw mu"></i>Fallecidos</span><span><i class="sw he"></i>Lesionados</span></div>
      <div class="graf-env" id="g-vict"></div></div>
    <div class="t-bloque"><h4>Comparendos por infracción en la línea base</h4><div class="cods" id="t-cods"></div></div>
  </div>
</section>

<section>
  <div><span class="eyebrow">Cifras a reportar</span><h2>Línea base por equipo</h2></div>
  <p class="lead">Totales de los 36 meses de cada equipo, tal como se cargan en la plataforma de la ANSV. El detalle mes a mes
    para copiar y pegar está en el archivo Excel adjunto. Un guion indica que el equipo no controla esa infracción.</p>
  <div class="tabla"><table><thead><tr><th>Equipo</th><th>Fallecidos</th><th>Lesionados</th>{cab}</tr></thead>
    <tbody>{"".join(filas)}</tbody></table></div>
  <p class="nota">{leyenda_cod}.</p>
</section>

<section>
  <div><span class="eyebrow">Método</span><h2>Cómo se construyó la línea base</h2></div>
  <ol class="pasos">
    <li><b>Zonas de influencia</b><p>Se tomó el área de influencia definida en el estudio técnico de cada equipo, con una tolerancia de 15 metros.</p></li>
    <li><b>Siniestros</b><p>Se ubicaron en el mapa los siniestros con víctimas del sistema de información de siniestralidad vial y se contaron las personas fallecidas y lesionadas.</p></li>
    <li><b>Comparendos</b><p>Se ubicaron los {fmt(c["n_comparendos_total"])} comparendos impuestos entre 2023 y 2026 a partir de la dirección de cada infracción.</p></li>
    <li><b>Conteo mensual</b><p>Se contó mes a mes en cada zona durante los 36 meses previos al inicio de operación de cada equipo.</p></li>
  </ol>
</section>

<section>
  <div><span class="eyebrow">Para tener en cuenta</span><h2>Lectura de las cifras</h2></div>
  <ul class="consideraciones">
    <li><b>Las cifras son un mínimo.</b> {c["pct_no_ubicable"]} de los comparendos no tenía una dirección suficientemente precisa para ubicarlo en el mapa y no se pudo asignar a ninguna zona.</li>
    <li><b>El registro de lesionados se fortaleció en 2026.</b> Los meses anteriores pueden mostrar menos lesionados de los que realmente hubo, por lo que conviene leerlos como un piso.</li>
    <li><b>Las cámaras no cuentan en su propia línea base.</b> Los comparendos generados por los equipos SAST desde su entrada en operación se registran aparte y servirán para el seguimiento.</li>
    <li><b>Cada equipo tiene su propio periodo.</b> La línea base cubre los tres años anteriores al mes en que cada equipo empezó a operar, según la plataforma de la ANSV.</li>
  </ul>
</section>

<section>
  <div><span class="eyebrow">Trazabilidad</span><h2>Fuentes, archivos y validación</h2></div>
  <div class="traza"><dl>{fuentes}</dl><dl>{entregables}</dl></div>
</section>

<footer class="pie"><span>Secretaría de Tránsito y Transporte de Valledupar</span><span>Generado el {esc(c["generado"])}</span>
  <span>Versión del cálculo <code>{esc(c["version"])}</code></span></footer>
</div>

<script src="https://cdnjs.cloudflare.com/ajax/libs/leaflet/1.9.4/leaflet.min.js"></script>
<script>
const D={datos};
const T={tablero};
const css=getComputedStyle(document.documentElement);const col=v=>css.getPropertyValue(v).trim();
const num=n=>Number(n).toLocaleString('es-CO');
const MES=['ene','feb','mar','abr','may','jun','jul','ago','sep','oct','nov','dic'];
const mesTxt=m=>MES[+m.slice(5)-1]+' '+m.slice(0,4);
function botonera(cont,items,alElegir){{
  cont.innerHTML='';const bs=[];
  items.forEach((it,i)=>{{const b=document.createElement('button');b.type='button';b.textContent=it.t;b.setAttribute('aria-pressed','false');
    b.onclick=()=>{{bs.forEach(x=>x.setAttribute('aria-pressed','false'));b.setAttribute('aria-pressed','true');alElegir(it,i);}};
    cont.appendChild(b);bs.push(b);}});
  return bs;
}}

// ---------------- mapa ----------------
(function(){{
  const el=document.getElementById('mapa');
  if(!window.L){{el.innerHTML='<p style="padding:20px">El mapa necesita conexión a internet.</p>';return;}}
  // la vista se fija antes de agregar capas: sin centro, Leaflet falla al dibujar los círculos
  const m=L.map(el,{{scrollWheelZoom:true}}).setView([10.47,-73.257],14);
  L.tileLayer('https://server.arcgisonline.com/ArcGIS/rest/services/World_Street_Map/MapServer/tile/{{z}}/{{y}}/{{x}}',
    {{maxZoom:19,attribution:'Fondo © Esri, HERE, Garmin, OpenStreetMap contributors'}}).addTo(m);
  const gZ=L.layerGroup().addTo(m),gC=L.layerGroup().addTo(m),gS=L.layerGroup().addTo(m),gE=L.layerGroup().addTo(m);
  const estado=document.getElementById('estado-mapa');
  function dibujar(sel,rotulo){{
    [gZ,gC,gS,gE].forEach(g=>g.clearLayers());
    const ok=l=>!sel||l.some(e=>sel.includes(e));
    L.geoJSON(D.zonas,{{filter:f=>!sel||sel.includes(f.properties.equipo),
      style:{{color:col('--accent'),weight:2,fillColor:col('--accent'),fillOpacity:.16}},
      onEachFeature:(f,l)=>l.bindTooltip('Zona de influencia '+f.properties.equipo)}}).addTo(gZ);
    let nc=0,ns=0;
    D.comparendos.forEach(p=>{{if(!ok(p[5]))return;nc++;
      L.circleMarker([p[1],p[0]],{{radius:4,weight:1,color:'#fff',fillColor:col('--comp'),fillOpacity:.85}})
        .bindTooltip(p[2]+' · '+p[3]+' · '+p[4]).addTo(gC);}});
    D.siniestros.forEach(p=>{{if(!ok(p[5]))return;ns++;const mu=p[3]>0;
      L.circleMarker([p[1],p[0]],{{radius:8,weight:2,color:'#fff',fillColor:mu?col('--dead'):col('--hurt'),fillOpacity:.95}})
        .bindTooltip('Siniestro del '+p[2]+'<br>'+(p[3]?p[3]+' fallecido'+(p[3]>1?'s':'')+'<br>':'')+(p[4]?p[4]+' lesionado'+(p[4]>1?'s':''):''))
        .addTo(gS);}});
    D.equipos.forEach(p=>{{const act=!sel||sel.includes(p[7]);
      L.marker([p[1],p[0]],{{icon:L.divIcon({{className:'placa-mapa',html:'<span>'+p[2]+'</span>',iconSize:[38,24]}}),
        zIndexOffset:act?1000:0,opacity:act?1:.35}})
        .bindTooltip('<b>Equipo '+p[2]+'</b> · '+p[3]+'<br>'+p[4]+'<br>Código único '+p[5]+' · inicio '+p[6]).addTo(gE);}});
    estado.textContent=rotulo+': '+num(nc)+' comparendos y '+ns+' siniestros con víctimas en la línea base.';
  }}
  const todo=()=>m.fitBounds(L.geoJSON(D.zonas).getBounds(),{{padding:[30,30]}});
  const cajaEq=document.getElementById('botones-eq');
  const puntos=[{{t:'Toda la ciudad'}}].concat(D.puntos.map(p=>({{t:p.nombre,p}})));
  const bs=botonera(document.getElementById('botones'),puntos,(it)=>{{
    if(!it.p){{cajaEq.hidden=true;dibujar(null,'Toda la ciudad');todo();return;}}
    const p=it.p;cajaEq.hidden=false;
    const sub=[{{t:'Todo el punto'}}].concat(p.equipos.map(e=>({{t:e,e}})));
    const be=botonera(cajaEq,sub,(x)=>{{
      if(!x.e){{dibujar(p.equipos,p.nombre);m.fitBounds(L.latLngBounds(p.limites),{{padding:[40,40],maxZoom:18}});return;}}
      dibujar([x.e],x.e+' · '+p.nombre);m.fitBounds(L.latLngBounds(D.limites_equipo[x.e]),{{padding:[50,50],maxZoom:19}});
    }});
    be[0].click();
  }});
  bs[0].click();
}})();

// ---------------- tablero ----------------
(function(){{
  const W=1000;
  function escala(mx){{const pasos=[1,2,5,10,20,25,50,100,200,250,500,1000];
    const p=pasos.find(x=>x>=mx/4)||1000;return {{paso:p,tope:Math.max(p,Math.ceil(mx/p)*p)}};}}
  function grafico(cont,u,series,H){{
    const meses=T.meses,n=meses.length,izq=46,der=10,arr=30,aba=28;
    const bw=(W-izq-der)/n,tot=meses.map((_,i)=>series.reduce((a,s)=>a+s.v[i],0));
    const {{paso,tope}}=escala(Math.max(1,...tot));const y=v=>arr+(H-arr-aba)*(1-v/tope);
    const i0=meses.indexOf(u.ventana[0]),i1=meses.indexOf(u.ventana[1]),ii=meses.indexOf(u.mes_inicio);
    let s='<svg viewBox="0 0 '+W+' '+H+'" role="img" aria-label="Serie mensual de '+u.nombre+'">';
    if(i0>=0&&i1>=0){{s+='<rect class="ventana" x="'+(izq+i0*bw)+'" y="'+arr+'" width="'+((i1-i0+1)*bw)+'" height="'+(H-arr-aba)+'"/>';
      for(let a=0;a<3;a++){{const xa=izq+(i0+a*12)*bw;if(a)s+='<line class="sepa" x1="'+xa+'" x2="'+xa+'" y1="'+(arr-4)+'" y2="'+(H-aba)+'"/>';
        s+='<text class="anio" x="'+(xa+6*bw)+'" y="'+(arr-10)+'" text-anchor="middle">Año '+(a+1)+'</text>';}}}}
    for(let k=0;k<=tope;k+=paso)s+='<line class="grid" x1="'+izq+'" x2="'+(W-der)+'" y1="'+y(k)+'" y2="'+y(k)+'"/>'+
      '<text class="eje" x="'+(izq-8)+'" y="'+(y(k)+4)+'" text-anchor="end">'+num(k)+'</text>';
    meses.forEach((m,i)=>{{const x=izq+i*bw;let base=0,g='<g class="col" data-i="'+i+'"><rect class="hit" x="'+x+'" y="'+arr+'" width="'+bw+'" height="'+(H-arr-aba)+'"/>';
      series.forEach(sr=>{{const v=sr.v[i];if(!v)return;const y0=y(base),y1=y(base+v);
        g+='<rect x="'+(x+1)+'" y="'+y1+'" width="'+Math.max(bw-2,1)+'" height="'+Math.max(y0-y1-(base?1.5:0),1)+'" style="fill:var('+sr.c+')"/>';base+=v;}});
      s+=g+'</g>';if(m.endsWith('-01'))s+='<text class="eje" x="'+(x+2)+'" y="'+(H-8)+'">'+m.slice(0,4)+'</text>';}});
    if(ii>=0){{const xi=izq+ii*bw;s+='<line class="inicio" x1="'+xi+'" x2="'+xi+'" y1="'+(arr-4)+'" y2="'+(H-aba)+'"/>'+
      '<text class="inicio-t" x="'+(xi-6)+'" y="'+(arr+12)+'" text-anchor="end">Inicio de operación</text>';}}
    s+='<line class="base" x1="'+izq+'" x2="'+(W-der)+'" y1="'+y(0)+'" y2="'+y(0)+'"/></svg><div class="tip" hidden></div>';
    cont.innerHTML=s;const tip=cont.querySelector('.tip'),svg=cont.querySelector('svg');
    cont.querySelectorAll('.col').forEach(gr=>{{
      gr.addEventListener('mouseenter',()=>{{const i=+gr.dataset.i,m=meses[i];
        const enBase=i>=i0&&i<=i1?'Línea base · Año '+(Math.floor((i-i0)/12)+1):(i>=ii&&ii>=0?'Equipo en operación':'Fuera de la línea base');
        tip.innerHTML='<b>'+mesTxt(m)+'</b> · '+enBase+'<br>'+series.map(sr=>'<i class="sw" style="background:var('+sr.c+')"></i>'+sr.n+': <b>'+num(sr.v[i])+'</b>').join('<br>')+
          (series.length>1?'<br>Total: <b>'+num(tot[i])+'</b>':'');
        tip.hidden=false;const r=svg.getBoundingClientRect(),esc=r.width/W;
        let left=(izq+(i+.5)*bw)*esc+12;if(left+tip.offsetWidth>r.width)left=(izq+(i+.5)*bw)*esc-tip.offsetWidth-12;
        tip.style.left=Math.max(0,left)+'px';tip.style.top='8px';}});
      gr.addEventListener('mouseleave',()=>{{tip.hidden=true;}});}});
  }}
  function codigos(u){{
    const mx=Math.max(1,...u.codigos.map(c=>c.agente+c.foto));
    document.getElementById('t-cods').innerHTML=u.codigos.slice().sort((a,b)=>(b.agente+b.foto)-(a.agente+a.foto)).map(c=>{{const t=c.agente+c.foto;
      return '<div class="cod'+(t?'':' cero')+'"><code>'+c.codigo+'</code><span class="desc">'+c.desc+'</span>'+
        '<span class="bar" title="Agentes '+num(c.agente)+' · Fotodetección '+num(c.foto)+'"><i class="ag" style="width:'+(100*c.agente/mx)+'%"></i><i class="fo" style="width:'+(100*c.foto/mx)+'%"></i></span>'+
        '<span class="tot">'+num(t)+'</span><span class="det">Agentes '+num(c.agente)+' · fotodetección '+num(c.foto)+'</span></div>';}}).join('');
  }}
  function mostrar(u,p){{
    document.getElementById('t-sol').textContent='Solicitud '+p.solicitud+(u===p?' · '+p.equipos.length+' equipos':' · '+p.nombre);
    document.getElementById('t-nombre').textContent=u===p?p.nombre:'Equipo '+u.nombre.slice(-3);
    document.getElementById('t-meta').textContent=(u===p?'Equipos '+p.equipos.map(e=>e.nombre.slice(-3)).join(' y ')+' · ':u.direccion+' · código único '+u.codigo_unico+' · ')+
      'inicio de operación '+u.inicio+' · línea base '+u.texto_ventana;
    document.getElementById('k-comp').textContent=num(u.tot.comparendos);
    document.getElementById('k-fall').textContent=num(u.tot.fallecidos);
    document.getElementById('k-les').textContent=num(u.tot.lesionados);
    grafico(document.getElementById('g-comp'),u,[{{n:'Agentes',c:'--s-ag',v:u.agente}},{{n:'Fotodetección anterior',c:'--s-fo',v:u.foto}}],290);
    grafico(document.getElementById('g-vict'),u,[{{n:'Fallecidos',c:'--dead',v:u.fallecidos}},{{n:'Lesionados',c:'--hurt',v:u.lesionados}}],150);
    codigos(u);
  }}
  const caja=document.getElementById('t-equipos');
  const bp=botonera(document.getElementById('t-puntos'),T.puntos.map(p=>({{t:p.nombre,p}})),(it)=>{{
    const p=it.p;const be=botonera(caja,[{{t:'Todo el punto',u:p}}].concat(p.equipos.map(e=>({{t:'Equipo '+e.nombre.slice(-3),u:e}}))),x=>mostrar(x.u,p));
    be[0].click();}});
  bp[0].click();
}})();
</script>
</body></html>"""
