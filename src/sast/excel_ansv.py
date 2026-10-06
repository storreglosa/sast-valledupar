"""Excel de línea base en el formato de la plataforma ANSV (una hoja por equipo).

Filas: Fallecidos, Lesionados y los códigos aprobados del equipo (en el orden del Excel de
equipos). Columnas: los 36 meses previos al inicio de operación del equipo agrupados en Año 1, 2
y 3, Total y
Observaciones. Valores: criterio oficial mixto (polígono para comparendos ubicados por dirección, +15 m para siniestros y GPS), comparendos de agentes + fotodetección
previa (sin SAST).
"""

from __future__ import annotations

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


MESES_ES = ["Enero", "Febrero", "Marzo", "Abril", "Mayo", "Junio", "Julio", "Agosto",
            "Septiembre", "Octubre", "Noviembre", "Diciembre"]
NARANJA = PatternFill("solid", fgColor="E07B4C")
GRIS = PatternFill("solid", fgColor="EEF0F2")
AZUL = PatternFill("solid", fgColor="DCE6F1")
FINO = Side(style="thin", color="BFBFBF")
BORDE = Border(left=FINO, right=FINO, top=FINO, bottom=FINO)
NEGRITA = Font(bold=True)
CENTRO = Alignment(horizontal="center", vertical="center", wrap_text=True)


def etiqueta_mes(p: pd.Period) -> str:
    return f"{MESES_ES[p.month - 1]} {p.year}"


def _hoja_tabla(ws, df: pd.DataFrame, fila: int = 1) -> None:
    for j, c in enumerate(df.columns, 1):
        cel = ws.cell(fila, j, str(c))
        cel.font, cel.fill, cel.border, cel.alignment = NEGRITA, GRIS, BORDE, CENTRO
    for i, r in enumerate(df.itertuples(index=False), fila + 1):
        for j, v in enumerate(r, 1):
            ws.cell(i, j, None if (isinstance(v, float) and pd.isna(v)) else v).border = BORDE
    for j, c in enumerate(df.columns, 1):
        ancho = max([len(str(c))] + [len(str(v)) for v in df.iloc[:, j - 1].head(200)])
        ws.column_dimensions[get_column_letter(j)].width = min(max(ancho + 2, 8), 60)
    ws.freeze_panes = ws.cell(fila + 1, 1)


def hoja_equipo(ws, info: dict, valores: pd.DataFrame, observaciones: dict) -> None:
    """valores: index = indicadores (orden de filas), columns = Period (36 meses)."""
    meses = list(info["ventana"])
    ws["A1"] = f"Línea base — {info['equipo']} — {info['punto']} (solicitud {info['solicitud']})"
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = f"Dirección: {info['direccion']}"
    ws["A3"] = (f"Fecha inicio de operación: {info['fecha_inicio']}    Código único: {info['codigo_unico']}    "
                f"Solicitud: {info['solicitud_ansv']}    Dirección ANSV: {info['direccion_ansv']}")
    ws["A4"] = (f"Criterio: zona de influencia (comparendos ubicados por dirección: polígono; siniestros y GPS: "
                f"+{info['buffer']:.0f} m, EPSG:9377). Comparendos de agentes "
                "y fotodetección previa (sin cámaras SAST). Fallecidos y lesionados = personas, portal ANSV "
                f"corte {info['corte_portal']}.")
    f0 = 6
    ws.cell(f0, 1, "Indicador")
    for a in range(3):
        c0 = 2 + a * 12
        ws.cell(f0, c0, f"Año {a + 1}")
        ws.merge_cells(start_row=f0, start_column=c0, end_row=f0, end_column=c0 + 11)
        ws.cell(f0, c0).fill = NARANJA if a == 0 else AZUL
        ws.cell(f0, c0).font = Font(bold=True, color="FFFFFF" if a == 0 else "000000")
        ws.cell(f0, c0).alignment = CENTRO
    ct, co = 2 + 36, 3 + 36
    ws.cell(f0 + 1, 1, "Indicador")
    for j, m in enumerate(meses):
        ws.cell(f0 + 1, 2 + j, etiqueta_mes(m))
    ws.cell(f0 + 1, ct, "Total")
    ws.cell(f0 + 1, co, "Observaciones")
    for j in range(1, co + 1):
        c = ws.cell(f0 + 1, j)
        c.font, c.fill, c.border, c.alignment = NEGRITA, GRIS, BORDE, CENTRO
    for i, ind in enumerate(valores.index):
        f = f0 + 2 + i
        ws.cell(f, 1, ind).font = NEGRITA
        ws.cell(f, 1).border = BORDE
        for j, m in enumerate(meses):
            c = ws.cell(f, 2 + j, int(valores.at[ind, m]))
            c.border, c.alignment = BORDE, CENTRO
        # valor fijo, no fórmula: un lector que no recalcula (pandas, plataforma web) vería la celda vacía
        c = ws.cell(f, ct, int(valores.loc[ind].sum()))
        c.font, c.border, c.alignment = NEGRITA, BORDE, CENTRO
        c = ws.cell(f, co, observaciones.get(ind, ""))
        c.border, c.alignment = BORDE, Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 13
    for j in range(2, 38):
        ws.column_dimensions[get_column_letter(j)].width = 10.5
    ws.column_dimensions[get_column_letter(co)].width = 70
    ws.row_dimensions[f0 + 1].height = 32
    ws.freeze_panes = ws.cell(f0 + 2, 2)


def libro(hojas_equipo: list[tuple[dict, pd.DataFrame, dict]], anexos: dict[str, pd.DataFrame],
          metodologia: list[str]) -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "Metodología"
    for i, t in enumerate(metodologia, 1):
        ws.cell(i, 1, t).alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 130
    for info, valores, obs in hojas_equipo:
        hoja_equipo(wb.create_sheet(info["equipo"]), info, valores, obs)
    for nombre, df in anexos.items():
        _hoja_tabla(wb.create_sheet(nombre[:31]), df)
    return wb
