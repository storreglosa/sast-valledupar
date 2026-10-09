"""Datos públicos del tablero (`tablero/data/*.json`).

Solo sale lo que está en las listas blancas de abajo: equipos SAST (públicos por la ANSV),
planes semafóricos, geometría OSM y cifras agregadas de línea base. Nada de registros puntuales.
La asignación grupo -> acceso solo se publica validada (decisión 28); con `borrador=True` se
escribe marcada como borrador para ensayos locales y `scripts/verificar_tablero.py` impide
publicarla.
"""

from __future__ import annotations

import json
import re

import openpyxl
import pandas as pd

from sast.ejecutivo import INFRACCIONES
from sast.semaforos.cruce import huella_config

ESQUEMAS = {"semaforos": "tablero-semaforos/1", "geometria": "tablero-geometria/1", "sast": "tablero-sast/1"}
NOTA_FASE = ("Fase ilustrativa: el plan es el que rige a esta hora, pero el segundo del ciclo no está "
             "sincronizado con el controlador.")
NOTA_VEHICULOS = "Vehículos y peatones ilustrativos: no son aforos."


def _txt(v):
    return None if v is None or (isinstance(v, float) and pd.isna(v)) else str(v)


def _mov(b: dict) -> dict:
    return {k: b.get(k) for k in ("codigo", "movimiento", "acceso", "sale_por", "brazo", "mitad", "via")}


def semaforos(proc: dict, conf: dict, borrador: bool) -> tuple[dict, dict]:
    """(semaforos.json, geometria.json) a partir de data/processed/semaforos.json y la config."""
    cfg = {c["id"]: c for c in conf["intersecciones"]}
    if proc.get("errores_sin_decision"):
        raise SystemExit(f"La etapa 9 terminó con {len(proc['errores_sin_decision'])} hallazgos de nivel ERROR sin "
                         "decisión (outputs/tables/semaforos_validacion.csv): no se publica")
    inters, geos = [], {}
    for it in proc["intersecciones"]:
        c = cfg[it["id"]]
        val = c.get("asignacion")
        if it.get("config_usada") != huella_config(c):
            raise SystemExit(f"{it['id']}: la config del cruce en config/semaforos.yaml (asignación, geometría, centro "
                             "o señalización) cambió desde la última corrida de la etapa 9; corre "
                             "scripts/09_semaforos.py antes de 10_tablero.py")
        if val and val.get("validado_por"):
            # validada: el borrador de la etapa 9 ya trae las correcciones (cambios) con su geometría
            mov = {g: {**_mov(b), "confirmada": True} for g, b in it["borrador"].items()}
            estado = "validada"
        elif borrador:
            estado, mov = "borrador", {g: _mov(b) for g, b in it["borrador"].items()}
        else:
            estado, mov = "pendiente", {}
        nombres = c.get("nombres_via") or {}
        for m in mov.values():
            if m.get("via") in nombres:
                m["via"] = nombres[m["via"]]
        publica = estado != "pendiente"
        grupos = [{"id": g["id"], "nombre": g["nombre"], "tipo": g["tipo"],
                   **({"mov": mov[g["id"]]} if g["id"] in mov else {})} for g in it["grupos"]]
        inters.append({
            "id": it["id"], "nombre": it["nombre"], "direccion": c["direccion"], "controlador": {"equipo": it["controlador"]["equipo"],
                                                                     "cruce": it["controlador"]["cruce"]},
            "cruce_pie": it["cruce_pie"].split("UBICADO EN ")[-1].strip(),
            "equipos": it["equipos"], "centro": {**it["geometria"]["origen"], "crs": "EPSG:4326"},
            "grupos": grupos, "amigos": it["amigos"],
            "planes": [{k: p[k] for k in ("id", "ciclo", "con_horario", "tiempos", "kpis", "etapas")}
                       for p in it["planes"]],
            "horario": it["horario"],
            "asignacion": {"estado": estado, **({"validado_por": val["validado_por"], "fecha": str(val["fecha"])}
                                                if estado == "validada" else {})},
        })
        g = it["geometria"]
        geos[it["id"]] = {
            "norte": g["norte_codificacion"] if publica else 0.0,
            "vias": [{"p": v["puntos"], "ancho": v["ancho"], "unico": v["unico"], "ctx": v["contexto"],
                      "n": (c.get("nombres_via") or {}).get(v["nomencla"] or "sin nombre", v["nomencla"])}
                     for v in g["vias"]],
            "brazos": [{k: b[k] for k in ("id", "rumbo", "nomencla", "r_caja") + (("cardinal",) if publica else ())}
                       for b in g["brazos"]],
            "cajon": g["cajon_amarillo"],
            "cebras": [{"grupo": k if k in mov else None, "codigo": mov.get(k, {}).get("codigo"),
                        "poligono": z["poligono"], "eje": z["eje"]}
                       for k, z in g["cebras"].items() if k.startswith("sin_semaforo") or (publica and k in mov)],
            "pare": {k: v for k, v in g["lineas_pare"].items() if publica and k in mov},
            "trayectorias": {k: v for k, v in g["trayectorias"].items() if publica and k in mov},
            "camaras": g["camaras"],
            # cruce largo (pares lejos del centro): encuadre del diagrama, solo si se dibujan los pares
            **({"vista": g["vista"]} if publica and g.get("vista") else {}),
        }
    sem = {"esquema": ESQUEMAS["semaforos"], "generado": proc["generado"], "zona_horaria": "America/Bogota",
           "nota_fase": NOTA_FASE, "nota_vehiculos": NOTA_VEHICULOS,
           "fuente": "Reportes del controlador SISTRA Wiseverse V3.0 (" + ", ".join(
               f"{f[8:10]}/{f[5:7]}/{f[:4]}" for f in sorted({a["archivo"].split("/")[-1][:10] for a in proc["fuente"]
                                                               if "_sistra_" in a["archivo"]}))
               + "), leídos por scripts/09_semaforos.py",
           "intersecciones": inters}
    geo = {"esquema": ESQUEMAS["geometria"], "unidad": "m", "crs_origen": "EPSG:9377 (CTM12), relativo al centro",
           "atribucion": "Vías © colaboradores de OpenStreetMap (ODbL), corte "
                         + next(a["archivo"].split("/")[-1][:10] for a in proc["fuente"] if a["archivo"].startswith("red_vial/")),
           "intersecciones": geos}
    return sem, geo


def _resumen_excel(ruta) -> tuple[pd.DataFrame, dict]:
    wb = openpyxl.load_workbook(ruta, read_only=True)
    filas = list(wb["Resumen"].iter_rows(values_only=True))
    res = pd.DataFrame(filas[1:], columns=filas[0]).dropna(subset=["Equipo"])
    texto = " ".join(str(r[0]) for r in wb["Instrucciones"].iter_rows(values_only=True) if r[0])
    cortes = {"siniestros": re.search(r"siniestralidad, corte (\d\d/\d\d/\d{4})", texto).group(1),
              "comparendos": re.search(r"Corte (\d\d/\d\d/\d{4})", texto).group(1),
              "version": re.search(r"Versión del cálculo: (\w+)", texto).group(1)}
    return res, cortes


def sast(equipos: pd.DataFrame, ruta_excel, largo: pd.DataFrame, geocod: pd.DataFrame, semaforo_de: dict) -> dict:
    """sast.json: cámaras (GeoJSON) y línea base oficial (+15 m, decisión 25) por equipo.

    Las cifras salen de la hoja «Resumen» del Excel entregado a la plataforma ANSV y se cruzan
    con outputs/tables/indicadores_largo.csv filtrada como 08_reporte_ejecutivo.py."""
    res, cortes = _resumen_excel(ruta_excel)
    l = largo[(largo["criterio"] == "buffer15") & largo["aprobado"].astype(bool)
              & largo["medio"].isin(["agente", "fotodeteccion_previa", "portal"])
              & largo["anio_base"].notna() & (largo["nivel"] == "equipo")]
    codigos = [c for c in res.columns if re.fullmatch(r"[A-Z]\d\d", str(c))]
    base, discrepan = {}, []
    for d in res.to_dict("records"):
        eq = d["Equipo"]
        porcod = {c: int(d[c]) for c in codigos if d[c] is not None and not pd.isna(d[c])}
        calc = l[l["unidad"] == eq].groupby("indicador")["valor"].sum()
        for k, v in [("Fallecidos", d["Fallecidos"]), ("Lesionados", d["Lesionados"]), *porcod.items()]:
            if int(calc.get(k, 0)) != int(v):
                discrepan.append(f"{eq} {k}: Excel {v}, tabla larga {int(calc.get(k, 0))}")
        medios = (l[(l["unidad"] == eq) & ~l["indicador"].isin(["Fallecidos", "Lesionados"])]
                  .groupby(["indicador", "medio"])["valor"].sum().unstack(fill_value=0))
        base[eq] = {
            "ventana": d["Línea base"], "inicio": d["Inicio de operación"],
            "fallecidos": int(d["Fallecidos"]), "lesionados": int(d["Lesionados"]),
            "comparendos": [{"codigo": c, "total": n,
                             "agente": int(medios.loc[c, "agente"]) if c in medios.index and "agente" in medios else 0,
                             "foto_previa": int(medios.loc[c, "fotodeteccion_previa"])
                             if c in medios.index and "fotodeteccion_previa" in medios else 0}
                            for c, n in sorted(porcod.items(), key=lambda x: -x[1])],
            "total_comparendos": sum(porcod.values()),
        }
    if discrepan:
        raise SystemExit("La línea base del Excel no cuadra con la tabla larga:\n  " + "\n  ".join(discrepan))

    g = geocod[geocod["medio"].isin(["agente", "fotodeteccion_previa"])]
    anios = [c for c in g.columns if c.isdigit()]
    pct = g[g["ubicacion"] == "no_ubicable"][anios].to_numpy().sum() / g[anios].to_numpy().sum()

    feats = []
    for r in equipos.itertuples():
        feats.append({"type": "Feature", "geometry": {"type": "Point", "coordinates": [round(r.lon, 6), round(r.lat, 6)]},
                      "properties": {"equipo": r.equipo, "numero": r.equipo[-3:], "solicitud": int(r.solicitud),
                                     "punto": r.punto.title(), "direccion": r.direccion,
                                     "direccion_ansv": _txt(r.direccion_ansv), "estado": r.estado,
                                     "fecha_inicio": None if pd.isna(r.fecha_inicio) else f"{r.fecha_inicio:%Y-%m-%d}",
                                     "codigo_unico": _txt(r.codigo_unico), "solicitud_ansv": _txt(r.solicitud_ansv),
                                     "codigos": [c.strip() for c in r.codigos.split(",") if c.strip()],
                                     "semaforo": semaforo_de.get(r.equipo)}})
    usados = sorted({c for f in feats for c in f["properties"]["codigos"]})
    return {
        "esquema": ESQUEMAS["sast"],
        "criterio": "Zona de influencia + 15 m: cifra oficial reportada a la ANSV (decisión 25).",
        "cortes": cortes,
        "salvedades": [
            f"Las cifras son un mínimo: el {pct:.0%} de los comparendos no tenía una dirección suficientemente "
            "precisa para ubicarlo y no se asignó a ninguna zona.",
            "El registro de lesionados se fortaleció en 2026: los meses anteriores pueden mostrar menos lesionados "
            "de los que hubo; léanse como un piso. Los fallecidos también cambian de fuente entre años.",
            "Los comparendos de las propias cámaras SAST no cuentan en su línea base; se informan aparte.",
            "Las zonas de las dos cámaras de un mismo punto pueden solaparse (p. ej. 021/022 y 041/042): un hecho "
            "cuenta para cada equipo (decisión 5), así que no se suman las cifras de dos equipos del mismo punto.",
            "Cada equipo tiene su periodo: los 36 meses anteriores al mes en que empezó a operar (plataforma ANSV).",
        ],
        "infracciones": {c: INFRACCIONES.get(c) for c in usados},
        "equipos": {"type": "FeatureCollection", "features": feats},
        "linea_base": base,
    }
